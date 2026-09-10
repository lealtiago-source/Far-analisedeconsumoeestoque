"""
Relatório de necessidade de compra.

Substitui analise_compra.py.

Correções relevantes de comportamento:

  1. FARMÁCIA POPULAR SEMPRE VAZIA — antes o teste era
     `medicamento.upper() in FARMACIA_POPULAR`, comparando o nome BRUTO da
     planilha contra chaves já normalizadas ("LOSARTANA 50 MG"). Como o nome
     bruto é algo como "LOSARTANA POTÁSSICA 50 MG COMPRIMIDO REVESTIDO", a
     comparação praticamente nunca casava e a seção 4 do relatório saía vazia.
     Agora ambos os lados passam por agrupar_medicamento().

  2. ACESSO A COLUNAS POR NOME — antes tudo era `row.iloc[4]`, `row.iloc[9]`
     etc. A planilha de entrada é gerada por analise.py, que remove e reordena
     colunas conforme a presença do arquivo de pedidos. Ou seja: as posições
     NÃO eram estáveis e um relatório podia ler validade na coluna de preço
     sem erro visível. Agora as colunas são localizadas pelo cabeçalho.

  3. DEDUPLICAÇÃO CONSISTENTE — `medicamentos_adicionados` usava o nome bruto
     enquanto os filtros finais usavam o nome normalizado. Grafias diferentes
     do mesmo item apareciam duas vezes. Agora tudo usa a chave normalizada.

  4. FALHA DE LICITAÇÃO VISÍVEL — cada aba tinha `except: print(...)`. Se
     todas falhassem, o relatório saía sem preços e sem aviso. Agora as falhas
     são registradas e devolvidas para exibição ao usuário.

  5. Coluna "Valor Total Estimado" acrescentada, e preços mantidos como número
     (antes o "-" textual misturado com floats tornava a coluna insomável).

  6. Os prints de depuração foram substituídos por logging.
"""

from __future__ import annotations

import logging
import math
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import (
    DIAS_ESTOQUE_PADRAO,
    DIAS_VALIDADE_PADRAO,
    MESES_COBERTURA,
    RESULTADO_DIR,
)
from normalizacao import agrupar_medicamento
from planilhas import PlanilhaInvalidaError, converter_numero, ler_planilha

logger = logging.getLogger(__name__)

# Chaves normalizadas uma única vez, no import.
FARMACIA_POPULAR = {
    agrupar_medicamento(nome)
    for nome in [
        "ALENDRONATO DE SODIO 70 MG",
        "ANLODIPINO 5 MG",
        "ATENOLOL 25 MG",
        "BECLOMETASONA 250 MCG/DOSE AEROSSOL ORAL",
        "BUDESONIDA 50 MCG NASAL",
        "ESPIRONOLACTONA 25 MG",
        "FUROSEMIDA 40 MG",
        "GLIBENCLAMIDA 5 MG",
        "HIDROCLOROTIAZIDA 25 MG",
        "LEVODOPA + BENSERAZIDA 100+25 MG",
        "PROPRANOLOL CLORIDRATO 40 MG",
        "ENALAPRIL 10 MG",
        "CAPTOPRIL 25 MG",
        "LOSARTANA 50 MG",
        "METFORMINA 500 MG",
        "METFORMINA 850 MG",
        "SALBUTAMOL 100 MCG AEROSSOL ORAL",
    ]
}

# Cabeçalhos esperados na planilha gerada por analise.py.
COL_MEDICAMENTO = "Medicamento"
COL_CONSUMO = "Consumo Mensal Médio Corrigido"
COL_ESTOQUE = "Quantidade em Estoque"
COL_VALIDADE = "Validade"
COL_LOTES = "Lotes/Validades Não Utilizados"
COL_PREVISAO = "Previsão Fim do Estoque"
COL_PEDIDOS = "Pedidos em Andamento"

# Abas da planilha de licitação: (nome da aba, coluna do nome, coluna do preço, linhas a pular)
ABAS_LICITACAO = [
    ("CISAMAPI 26", 1, 7, 0),
    ("Sigaf MED I", 0, 7, 8),
    ("Sigaf MED IV", 0, 6, 8),
    ("Prefeitura", 4, 5, 0),
]

VALORES_IGNORADOS = {"", "nan", "none", "total", "medicamento", "produto"}


@dataclass
class Item:
    medicamento: str
    motivo: str
    detalhe: str
    quantidade: int
    preco: float | None = None
    origem: str = ""

    @property
    def valor_total(self) -> float | None:
        return round(self.preco * self.quantidade, 2) if self.preco is not None else None


@dataclass
class ResultadoCompra:
    caminho_word: Path
    caminho_excel: Path
    falta: list[Item] = field(default_factory=list)
    validade_critica: list[Item] = field(default_factory=list)
    previsao_estoque: list[Item] = field(default_factory=list)
    farmacia_popular: list[Item] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Licitação
# ---------------------------------------------------------------------------


def carregar_dados_licitacao(
    caminho: str | Path | None,
) -> tuple[dict[str, dict], list[str]]:
    """
    Devolve o menor preço por medicamento normalizado e a lista de avisos.

    Retorno: ({chave: {"preco": float, "aba": str, "original": str}}, avisos)
    """
    avisos: list[str] = []
    if not caminho or not Path(caminho).exists():
        return {}, avisos

    linhas: list[dict] = []

    for aba, col_nome, col_preco, skiprows in ABAS_LICITACAO:
        try:
            df = pd.read_excel(caminho, sheet_name=aba, skiprows=skiprows, header=0)
        except Exception as exc:  # aba ausente ou ilegível
            aviso = f"Aba de licitação '{aba}' não foi lida: {exc}"
            logger.warning(aviso)
            avisos.append(aviso)
            continue

        if df.shape[1] <= max(col_nome, col_preco):
            aviso = (
                f"Aba '{aba}' tem {df.shape[1]} colunas; eram esperadas ao menos "
                f"{max(col_nome, col_preco) + 1}. Aba ignorada."
            )
            logger.warning(aviso)
            avisos.append(aviso)
            continue

        for _, linha in df.iterrows():
            nome_original = linha.iloc[col_nome]
            if pd.isna(nome_original):
                continue

            preco = converter_numero(linha.iloc[col_preco])
            if preco <= 0:
                continue

            chave = agrupar_medicamento(str(nome_original))
            if not chave:
                continue

            linhas.append(
                {
                    "Aba": aba,
                    "Original": str(nome_original).strip(),
                    "Normalizado": chave,
                    "Preco": preco,
                }
            )

    if not linhas:
        avisos.append(
            "Nenhum preço válido foi extraído do arquivo de licitação. "
            "O relatório será gerado sem valores."
        )
        return {}, avisos

    df_licitacao = pd.DataFrame(linhas)

    mapeamento: dict[str, dict] = {}
    for chave, grupo in df_licitacao.groupby("Normalizado", sort=False):
        menor = grupo.loc[grupo["Preco"].idxmin()]
        mapeamento[chave] = {
            "preco": float(menor["Preco"]),
            "aba": str(menor["Aba"]),
            "original": str(menor["Original"]),
        }

    logger.info("%s medicamentos com preço carregados da licitação.", len(mapeamento))
    return mapeamento, avisos


# ---------------------------------------------------------------------------
# Validades
# ---------------------------------------------------------------------------

import re

_DATA_RE = re.compile(r"(\d{2}/\d{2}/\d{4})")


def obter_validade_mais_distante(validade: object, texto_lotes: object) -> datetime | None:
    """Considera a validade principal e as datas embutidas no texto de lotes."""
    datas: list[pd.Timestamp] = []

    principal = pd.to_datetime(validade, errors="coerce")
    if pd.notna(principal):
        datas.append(principal)

    if pd.notna(texto_lotes):
        for encontrada in _DATA_RE.findall(str(texto_lotes)):
            convertida = pd.to_datetime(encontrada, format="%d/%m/%Y", errors="coerce")
            if pd.notna(convertida):
                datas.append(convertida)

    return max(datas).to_pydatetime() if datas else None


# ---------------------------------------------------------------------------
# Documento Word
# ---------------------------------------------------------------------------


def _adicionar_tabela(doc: Document, cabecalhos: list[str], linhas: list[list]) -> None:
    tabela = doc.add_table(rows=1, cols=len(cabecalhos))
    tabela.style = "Table Grid"

    for i, titulo in enumerate(cabecalhos):
        celula = tabela.rows[0].cells[i]
        celula.text = str(titulo)
        for run in celula.paragraphs[0].runs:
            run.bold = True

    for linha in linhas:
        celulas = tabela.add_row().cells
        for i, valor in enumerate(linha):
            celulas[i].text = "" if valor is None else str(valor)

    doc.add_paragraph()


def _formatar_moeda(valor: float | None) -> str:
    if valor is None:
        return "-"
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _linhas_tabela(itens: list[Item], com_preco: bool, rotulo_detalhe: str) -> list[list]:
    linhas = []
    for item in itens:
        linha = [item.medicamento, item.detalhe, item.quantidade]
        if com_preco:
            linha.extend([_formatar_moeda(item.preco), item.origem, _formatar_moeda(item.valor_total)])
        linhas.append(linha)
    return linhas


# ---------------------------------------------------------------------------
# Análise principal
# ---------------------------------------------------------------------------


def executar_analise_compra(
    arquivo_analise: str | Path,
    dias_validade: int = DIAS_VALIDADE_PADRAO,
    dias_estoque: int = DIAS_ESTOQUE_PADRAO,
    arquivo_licitacao: str | Path | None = None,
    meses_cobertura: int = MESES_COBERTURA,
) -> ResultadoCompra:
    """Gera o relatório Word + planilha consolidada de necessidade de compra."""
    df = ler_planilha(arquivo_analise)

    obrigatorias = [COL_MEDICAMENTO, COL_CONSUMO, COL_ESTOQUE, COL_PREVISAO]
    faltando = [c for c in obrigatorias if c not in df.columns]
    if faltando:
        raise PlanilhaInvalidaError(
            "A planilha de análise não contém a(s) coluna(s): "
            f"{', '.join(faltando)}. Gere-a novamente pela Análise Principal."
        )

    tem_pedidos = COL_PEDIDOS in df.columns
    df["__chave__"] = df[COL_MEDICAMENTO].map(agrupar_medicamento)

    # Coleta as chaves com pedido em andamento ANTES de filtrar.
    if tem_pedidos:
        com_pedido = set(
            df.loc[
                df[COL_PEDIDOS].fillna("").astype(str).str.strip() != "", "__chave__"
            ]
        )
        df = df[df[COL_PEDIDOS].fillna("").astype(str).str.strip() == ""]
    else:
        com_pedido = set()

    logger.info("%s medicamentos excluídos por pedido em andamento.", len(com_pedido))

    hoje = datetime.today()
    limite_validade = hoje + timedelta(days=dias_validade)
    limite_estoque = hoje + timedelta(days=dias_estoque)

    precos, avisos = carregar_dados_licitacao(arquivo_licitacao)
    com_preco = bool(arquivo_licitacao)

    resultado = ResultadoCompra(
        caminho_word=RESULTADO_DIR, caminho_excel=RESULTADO_DIR, avisos=avisos
    )
    vistos: set[str] = set()

    for _, linha in df.iterrows():
        medicamento = str(linha[COL_MEDICAMENTO]).strip()
        chave = linha["__chave__"]

        if not chave or medicamento.lower() in VALORES_IGNORADOS:
            continue
        if chave in com_pedido or chave in vistos:
            continue

        consumo = converter_numero(linha[COL_CONSUMO])
        quantidade = math.ceil(consumo * meses_cobertura)
        if quantidade <= 0:
            continue

        dados_preco = precos.get(chave)
        preco = dados_preco["preco"] if dados_preco else None
        origem = dados_preco["aba"] if dados_preco else ("Não encontrado" if com_preco else "")

        eh_farmacia_popular = chave in FARMACIA_POPULAR

        estoque_txt = str(linha[COL_ESTOQUE]).upper()
        validade = obter_validade_mais_distante(
            linha.get(COL_VALIDADE), linha.get(COL_LOTES)
        )
        previsao = pd.to_datetime(linha[COL_PREVISAO], errors="coerce")

        if "FALTA" in estoque_txt:
            motivo, detalhe, destino = "FALTA", str(linha[COL_ESTOQUE]), resultado.falta
        elif validade is not None and validade <= limite_validade:
            motivo = "VALIDADE CRÍTICA"
            detalhe = validade.strftime("%d/%m/%Y")
            destino = resultado.validade_critica
        elif pd.notna(previsao) and previsao <= limite_estoque:
            motivo = "PREVISÃO DE ESTOQUE"
            detalhe = previsao.strftime("%d/%m/%Y")
            destino = resultado.previsao_estoque
        else:
            continue

        item = Item(
            medicamento=medicamento,
            motivo=motivo,
            detalhe=motivo if eh_farmacia_popular else detalhe,
            quantidade=quantidade,
            preco=preco,
            origem=origem,
        )

        if eh_farmacia_popular:
            resultado.farmacia_popular.append(item)
        else:
            destino.append(item)

        vistos.add(chave)

    # ---------------- Word ----------------

    doc = Document()
    titulo = doc.add_heading("RELATÓRIO DE NECESSIDADE DE COMPRA DE MEDICAMENTOS", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Data de emissão: {hoje.strftime('%d/%m/%Y')}")
    doc.add_paragraph(
        "Parâmetros utilizados:\n"
        f"• Validade crítica: {dias_validade} dias\n"
        f"• Previsão de término do estoque: {dias_estoque} dias\n"
        f"• Cobertura da compra: {meses_cobertura} meses"
    )
    doc.add_paragraph(
        "Relatório elaborado com base na análise de consumo, estoque, validade e "
        "previsão de abastecimento. Itens com pedido de compra em andamento foram "
        "excluídos automaticamente."
    )

    if resultado.avisos:
        doc.add_heading("Avisos de processamento", level=2)
        for aviso in resultado.avisos:
            doc.add_paragraph(aviso, style="List Bullet")

    extras = ["Menor Preço", "Origem/Aba", "Valor Total Estimado"] if com_preco else []

    secoes = [
        ("1. ITENS EM FALTA", "Falta desde", resultado.falta),
        ("2. ITENS COM VALIDADE CRÍTICA", "Validade Considerada", resultado.validade_critica),
        ("3. ITENS COM PREVISÃO DE TÉRMINO DE ESTOQUE", "Previsão Fim do Estoque", resultado.previsao_estoque),
        ("4. MEDICAMENTOS DA FARMÁCIA POPULAR", "Motivo", resultado.farmacia_popular),
    ]

    for titulo_secao, rotulo, itens in secoes:
        doc.add_heading(f"{titulo_secao} ({len(itens)})", level=2)
        cabecalho = [
            "Medicamento",
            rotulo,
            f"Quantidade para {meses_cobertura} meses",
        ] + extras
        _adicionar_tabela(doc, cabecalho, _linhas_tabela(itens, com_preco, rotulo))

    todos = (
        resultado.falta
        + resultado.validade_critica
        + resultado.previsao_estoque
        + resultado.farmacia_popular
    )
    valor_estimado = sum(i.valor_total or 0 for i in todos)

    doc.add_heading("5. RESUMO", level=2)
    _adicionar_tabela(
        doc,
        ["Indicador", "Quantidade"],
        [
            ["Itens em Falta", len(resultado.falta)],
            ["Itens com Validade Crítica", len(resultado.validade_critica)],
            ["Itens com Previsão de Término", len(resultado.previsao_estoque)],
            ["Itens Farmácia Popular", len(resultado.farmacia_popular)],
            ["Total de Itens para Compra", len(todos)],
            ["Valor Total Estimado", _formatar_moeda(valor_estimado) if com_preco else "-"],
        ],
    )

    # ---------------- Excel ----------------

    consolidado = [
        {
            "Medicamento": item.medicamento,
            "Motivo": item.motivo,
            "Quantidade a Comprar": item.quantidade,
            **(
                {
                    "Menor Preço": item.preco,
                    "Origem/Aba": item.origem,
                    "Valor Total Estimado": item.valor_total,
                }
                if com_preco
                else {}
            ),
        }
        for item in sorted(todos, key=lambda i: i.medicamento)
    ]

    data_relatorio = hoje.strftime("%d-%m-%Y")
    resultado.caminho_word = RESULTADO_DIR / f"Relatorio_Compra_Medicamentos_{data_relatorio}.docx"
    resultado.caminho_excel = RESULTADO_DIR / f"Relatorio_Compra_Medicamentos_{data_relatorio}.xlsx"

    doc.save(resultado.caminho_word)
    pd.DataFrame(consolidado).to_excel(resultado.caminho_excel, index=False)

    logger.info("Relatório de compra gerado com %s itens.", len(todos))
    return resultado
