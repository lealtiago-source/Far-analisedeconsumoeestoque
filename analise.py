"""
Análise de consumo, estoque e previsão de abastecimento.

Substitui analise.py.

Correções relevantes de comportamento:

  1. ERRO SILENCIOSO — antes a função terminava em
     `except Exception as e: print(f"Erro geral: {e}")`. Qualquer falha era
     impressa no console e a função retornava normalmente, então a rota Flask
     exibia "✅ Análise concluída" e oferecia para download o resultado da
     execução ANTERIOR. Agora as exceções propagam.

  2. PERÍODO POR MEDICAMENTO — `dias_totais` era o intervalo global do
     dataset. Um item que passou a ser dispensado só no último mês tinha o
     consumo diluído por todo o período e era subestimado. Agora o período
     é medido por medicamento.

  3. TETO NO DESCONTO DE FALTA — `dias_validos = max(dias - falta, 1)` podia
     chegar a 1 dia, e `consumo/1*30` inflava o consumo mensal em até 30x,
     estourando a quantidade de compra. Agora há piso proporcional.

  4. NORMALIZAÇÃO ÚNICA — `agrupar_medicamento` era aplicado duas vezes em
     cada coluna. É a função mais cara do sistema; agora roda uma vez (e com
     cache).

  5. Datas inválidas (NaT) no cálculo de lotes não utilizados deixaram de
     estourar exceção dentro do join.
"""

from __future__ import annotations

import logging
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from config import (
    DIAS_MINIMOS_DESABASTECIMENTO,
    DIAS_MINIMOS_PERIODO,
    PROPORCAO_MINIMA_PERIODO_VALIDO,
    RESULTADO_DIR,
)
from consumo import calcular_consumo
from normalizacao import agrupar_medicamento
from planilhas import PlanilhaInvalidaError, converter_datas, exigir_colunas, ler_planilha

logger = logging.getLogger(__name__)

FILL_VERMELHO = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
FILL_ROXO = PatternFill(start_color="FFB4A7D6", end_color="FFB4A7D6", fill_type="solid")
FILL_VERDE = PatternFill(start_color="FFA9D08E", end_color="FFA9D08E", fill_type="solid")
FILL_AMARELO = PatternFill(start_color="FFFFD966", end_color="FFFFD966", fill_type="solid")

COL_DISPENSACAO = ["Data Dispensação", "Medicamento/Produto", "Lote", "Quantidade Dispensada"]
COL_DISTRIBUICAO = [
    "Data Distribuição",
    "Medicamento/Produto",
    "Lote",
    "Quantidade distribuída (unidades)",
]
COL_ESTOQUE = ["Medicamento/Produto", "Lote", "Quantidade em Estoque", "Validade"]

# Posições usadas na planilha de pedidos (XML). Documentadas explicitamente
# porque a origem não tem cabeçalho estável.
XML_COL_MEDICAMENTO = 4
XML_COL_QUANTIDADE = 5
XML_COL_DATA = 8
XML_COL_EMPENHO = 9
XML_COL_ENTREGA = 11


def _consumo_por_medicamento(df: pd.DataFrame) -> pd.DataFrame:
    """
    Consumo mensal médio corrigido por dias de desabastecimento.

    A implementação vive em consumo.py porque analise.py, analise_compra.py e
    analise_orcamento.py mantinham três cópias divergentes do mesmo cálculo.
    """
    return calcular_consumo(df)



def _lotes_nao_utilizados(df_estoque: pd.DataFrame) -> dict[str, str]:
    """Pré-calcula o texto de lotes secundários (antes: O(n²) linha a linha)."""
    resultado: dict[str, str] = {}

    for medicamento, dados in df_estoque.groupby("Medicamento Agrupado", sort=False):
        if len(dados) <= 1:
            resultado[medicamento] = ""
            continue

        validade_min = dados["Validade"].min()
        outros = dados[dados["Validade"] > validade_min]

        partes = []
        for _, linha in outros.iterrows():
            validade = linha["Validade"]
            # NaT deixou de estourar dentro do join.
            texto_validade = (
                validade.strftime("%d/%m/%Y") if pd.notna(validade) else "sem validade"
            )
            partes.append(f"{linha['Lote']} ({texto_validade})")

        resultado[medicamento] = "; ".join(partes)

    return resultado


def _pedidos_em_andamento(caminho_xml: str | Path) -> dict[str, str]:
    """Agrupa os pedidos pendentes (sem data de entrega) por medicamento."""
    from planilhas import converter_quantidade

    df = ler_planilha(caminho_xml)

    if df.shape[1] <= XML_COL_ENTREGA:
        raise PlanilhaInvalidaError(
            f"A planilha de pedidos tem {df.shape[1]} colunas; "
            f"são esperadas ao menos {XML_COL_ENTREGA + 1}."
        )

    entrega = df.iloc[:, XML_COL_ENTREGA].fillna("").astype(str).str.strip()
    pendentes = df[entrega == ""].copy()
    pendentes["Medicamento Agrupado"] = pendentes.iloc[:, XML_COL_MEDICAMENTO].map(
        agrupar_medicamento
    )

    resultado: dict[str, str] = {}
    for medicamento, dados in pendentes.groupby("Medicamento Agrupado", sort=False):
        itens = []
        for _, linha in dados.iterrows():
            quantidade = converter_quantidade(linha.iloc[XML_COL_QUANTIDADE])
            empenho = linha.iloc[XML_COL_EMPENHO]
            empenho = empenho if pd.notna(empenho) else "SEM EMPENHO"
            data = pd.to_datetime(linha.iloc[XML_COL_DATA], errors="coerce")
            data_txt = data.strftime("%d/%m/%Y") if pd.notna(data) else "SEM DATA"
            itens.append(f"Qtd:{quantidade} | Emp:{empenho} | Data:{data_txt}")
        resultado[medicamento] = " ; ".join(itens)

    return resultado


def _aplicar_cores(caminho: Path, dias_alerta: int = 120) -> None:
    wb = load_workbook(caminho)
    ws = wb.active

    colunas = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    col_prev = colunas.get("Previsão Fim do Estoque")
    col_val = colunas.get("Validade")

    if not col_prev or not col_val:
        wb.save(caminho)
        return

    limite = datetime.today() + timedelta(days=dias_alerta)
    margem = timedelta(days=dias_alerta)

    for linha in range(2, ws.max_row + 1):
        previsao = pd.to_datetime(ws.cell(linha, col_prev).value, errors="coerce")
        validade = pd.to_datetime(ws.cell(linha, col_val).value, errors="coerce")

        ws.cell(linha, col_prev).number_format = "DD/MM/YYYY"
        ws.cell(linha, col_val).number_format = "DD/MM/YYYY"

        # Antes: comparações com NaT dentro de `try/except: continue`,
        # o que escondia linhas sem cor sem qualquer indicação.
        if pd.isna(previsao) or pd.isna(validade):
            continue

        if validade < previsao - margem:
            ws.cell(linha, col_val).fill = FILL_ROXO
        elif validade < previsao:
            ws.cell(linha, col_val).fill = FILL_AMARELO
        else:
            ws.cell(linha, col_val).fill = FILL_VERDE

        if previsao < limite:
            ws.cell(linha, col_prev).fill = FILL_VERMELHO

    wb.save(caminho)


def executar_analise(
    arquivo_dispensacao: str | Path,
    arquivo_distribuicao: str | Path,
    arquivo_estoque: str | Path,
    arquivo_xml: str | Path | None = None,
    nome_saida: str = "resultado_analise.xlsx",
) -> Path:
    """
    Gera a planilha de análise consolidada e devolve o caminho do arquivo.

    Levanta PlanilhaInvalidaError quando alguma planilha de entrada não tem o
    formato esperado. NÃO engole exceções: a rota que chama precisa saber que
    a análise falhou.
    """
    df_disp = ler_planilha(arquivo_dispensacao)
    exigir_colunas(df_disp, COL_DISPENSACAO, "dispensação")
    df_disp = df_disp[COL_DISPENSACAO].rename(
        columns={"Data Dispensação": "Data", "Quantidade Dispensada": "Quantidade"}
    )

    df_dist = ler_planilha(arquivo_distribuicao)
    exigir_colunas(df_dist, COL_DISTRIBUICAO, "distribuição")
    df_dist = df_dist[COL_DISTRIBUICAO].rename(
        columns={
            "Data Distribuição": "Data",
            "Quantidade distribuída (unidades)": "Quantidade",
        }
    )

    df_estoque = ler_planilha(arquivo_estoque)
    exigir_colunas(df_estoque, COL_ESTOQUE, "estoque")

    df_movimento = pd.concat([df_disp, df_dist], ignore_index=True)

    # Uma única normalização por coluna (antes: duas).
    df_movimento["Medicamento Agrupado"] = df_movimento["Medicamento/Produto"].map(
        agrupar_medicamento
    )
    df_estoque["Medicamento Agrupado"] = df_estoque["Medicamento/Produto"].map(
        agrupar_medicamento
    )

    df_movimento["Data"] = converter_datas(df_movimento["Data"], "movimentação")
    df_movimento["Quantidade"] = pd.to_numeric(
        df_movimento["Quantidade"], errors="coerce"
    ).fillna(0)
    df_estoque["Validade"] = converter_datas(df_estoque["Validade"], "estoque (validade)")
    df_estoque["Quantidade em Estoque"] = pd.to_numeric(
        df_estoque["Quantidade em Estoque"], errors="coerce"
    ).fillna(0)

    resumo = _consumo_por_medicamento(df_movimento)
    if resumo.empty:
        raise PlanilhaInvalidaError(
            "Nenhuma movimentação válida encontrada nas planilhas de "
            "dispensação e distribuição."
        )

    estoque_agrupado = (
        df_estoque.groupby("Medicamento Agrupado")
        .agg({"Quantidade em Estoque": "sum", "Validade": "min"})
        .reset_index()
    )
    resumo = resumo.merge(estoque_agrupado, on="Medicamento Agrupado", how="left")

    mapa_lotes = _lotes_nao_utilizados(df_estoque)
    resumo["Lotes/Validades Não Utilizados"] = (
        resumo["Medicamento Agrupado"].map(mapa_lotes).fillna("")
    )

    if arquivo_xml:
        mapa_pedidos = _pedidos_em_andamento(arquivo_xml)
        resumo["Pedidos em Andamento"] = (
            resumo["Medicamento Agrupado"].map(mapa_pedidos).fillna("")
        )
    else:
        resumo["Pedidos em Andamento"] = ""

    def prever(linha: pd.Series) -> datetime | None:
        estoque = float(linha["Quantidade em Estoque"] or 0)
        consumo = float(linha["Consumo Mensal Médio Corrigido"] or 0)
        if estoque > 0 and consumo > 0:
            return datetime.today() + timedelta(days=(estoque / consumo) * 30)
        return None

    resumo["Previsão Fim do Estoque"] = resumo.apply(prever, axis=1)

    def rotular_estoque(linha: pd.Series) -> object:
        quantidade = linha["Quantidade em Estoque"]
        if pd.isna(quantidade) or quantidade == 0:
            ultima = linha["Última Dispensação"]
            if pd.notna(ultima):
                return f"FALTA DESDE {ultima.strftime('%d/%m/%Y')}"
            return "FALTA"
        return quantidade

    resumo["Quantidade em Estoque"] = resumo.apply(rotular_estoque, axis=1)

    final = resumo.drop(columns=["Medicamento Agrupado", "Última Dispensação"])
    ordem = [c for c in final.columns if c != "Pedidos em Andamento"]
    ordem.append("Pedidos em Andamento")
    final = final[ordem]

    caminho_saida = RESULTADO_DIR / nome_saida
    final.to_excel(caminho_saida, index=False)
    _aplicar_cores(caminho_saida)

    logger.info("Análise concluída: %s linhas em %s", len(final), caminho_saida)
    return caminho_saida
