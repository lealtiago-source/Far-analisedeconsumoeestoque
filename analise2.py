import pandas as pd
from datetime import datetime
import os
from openpyxl.styles import Font
from openpyxl import load_workbook

def executar_analise_remume(caminho_estoque, caminho_correspondencias):
    df_estoque = pd.read_excel(caminho_estoque)
    df_mapeamento = pd.read_excel(caminho_correspondencias)

    col_remume = 'RELAÇÃO MUNICIPAL DE MEDICAMENTOS ESSENCIAIS'
    col_estoque = 'Medicamento/Produto'

    # Normalizar nomes
    df_estoque['Medicamento/Produto'] = df_estoque['Medicamento/Produto'].astype(str).str.lower().str.strip()
    df_mapeamento[col_estoque] = df_mapeamento[col_estoque].astype(str).str.lower().str.strip()
    df_mapeamento[col_remume] = df_mapeamento[col_remume].astype(str).str.lower().str.strip()

    resultado = []

    for _, row in df_mapeamento.iterrows():
        nome_remume = row[col_remume]
        nome_estoque = row[col_estoque]

        correspondentes = df_estoque[df_estoque['Medicamento/Produto'] == nome_estoque]

        if not correspondentes.empty:
            total_estoque = correspondentes['Quantidade em Estoque'].sum()
        else:
            total_estoque = 'EM FALTA'

        resultado.append({
            'Medicamento REMUME': nome_remume,
            'Medicamento no Estoque': nome_estoque,
            'Quantidade em Estoque': total_estoque
        })

    df_resultado = pd.DataFrame(resultado)

    # Exportar Excel
    hoje = datetime.today().strftime('%d-%m-%Y')
    nome_arquivo = f'Estoque_REMUME_MAPEADO_{hoje}.xlsx'
    caminho_saida = os.path.join('static', nome_arquivo)

    with pd.ExcelWriter(caminho_saida, engine='openpyxl') as writer:
        df_resultado.to_excel(writer, index=False, sheet_name='REMUME')
        ws = writer.sheets['REMUME']
        ws.insert_rows(1)
        ws.merge_cells('A1:C1')
        cell = ws['A1']
        cell.value = f'Estoque REMUME com correspondência mapeada - {hoje}'
        cell.font = Font(bold=True)

    return nome_arquivo
