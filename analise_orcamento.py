"""
Análise de orçamento — projeção anual de custo por medicamento.

Substitui analise_orcamento.py.

Problemas do original corrigidos aqui:

  1. PERÍODO GLOBAL — `dias_totais` era o intervalo de TODO o dataset e era
     aplicado a cada medicamento. Item que começou a circular no último mês
     tinha o custo anual subestimado; item descontinuado, superestimado.

  2. FALTA SEM TETO — `dias_validos = max(dias_totais - dias_falta, 1)`.
     Com um único par de lotes distantes, o divisor caía para 1 dia e o
     custo anual previsto era multiplicado por até 30. Em orçamento isso
     significa pedir verba errada. Agora há piso proporcional (config).

  3. `dias_falta += diff` somava o intervalo INTEIRO como desabastecimento,
     inclusive os 15 dias considerados normais de reposição.

  4. MÉDIA SIMPLES DE PREÇO — `dados['Valor Unitário'].mean()` dava o mesmo
     peso a um lote de 10 unidades e a um de 10.000. Agora é ponderada pela
     quantidade.

  5. SELEÇÃO DE COLUNAS SEM VALIDAÇÃO — `df[[...]]` com nome divergente
     levantava KeyError, capturado pelo `except Exception` que apenas dava
     print e re-levantava; a rota Flask não distinguia isso de erro real.
     Agora `exigir_colunas` produz mensagem legível para o usuário.

  6. VALORES EM TEXTO — "1.234,56" vindo do relatório virava NaN e o custo
     ia a zero em silêncio. Agora passa por `converter_numero`.

  7. GRAVAÇÃO EM `static/` com nome fixo: dois usuários simultâneos
     sobrescreviam o resultado um do outro e o arquivo ficava público.
     Agora vai para RESULTADO_DIR com data/hora no nome.

  8. `except: print(...)` — a mensagem ia para um console que ninguém vê no
     modo empacotado. Agora usa logging e as exceções propagam.
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import RESULTADO_DIR
from consumo import calcular_consumo
from normalizacao import agrupar_medicamento
from planilhas import converter_datas, converter_numero, exigir_colunas, ler_planilha

logger = logging.getLogger(__name__)

FILL_VERDE = PatternFill(start_color="FFA9D08E", end_color="FFA9D08E", fill_type="solid")
FILL_AMARELO = PatternFill(start_color="FFFFD966", end_color="FFFFD966", fill_type="solid")
FILL_VERMELHO = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
FILL_CABECALHO = PatternFill(start_color="FF305496", end_color="FF305496", fill_type="solid")

COL_DISPENSACAO = [
    "Data Dispensação",
    "Medicamento/Produto",
    "Lote",
    "Quantidade Dispensada",
    "Valor Unitário",
]
COL_DISTRIBUICAO = [
    "Data Distribuição",
    "Medicamento/Produto",
    "Lote",
    "Quantidade distribuída (unidades)",
    "Valor unitário",
]

# Faixas de destaque do custo anual (antes hardcoded no meio do loop).
LIMITE_CUSTO_ALTO = 100_000.0
LIMITE_CUSTO_MEDIO = 50_000.0


def _preparar(caminho: Path | str, colunas: list[str], origem: str, fonte: str) -> pd.DataFrame:
    df = ler_planilha(caminho)
    exigir_colunas(df, colunas, origem)

    df = df[colunas].copy()
    df.columns = ["Data", "Medicamento/Produto", "Lote", "Quantidade", "Valor Unitário"]

    df["Fonte"] = fonte
    df["Data"] = converter_datas(df["Data"], origem)
    df["Quantidade"] = df["Quantidade"].map(converter_numero)
    df["Valor Unitário"] = df["Valor Unitário"].map(converter_numero)
    df["Lote"] = df["Lote"].astype(str).str.strip()

    return df


def _formatar(caminho: Path) -> None:
    wb = load_workbook(caminho)
    ws = wb.active

    cabecalhos = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}

    for coluna in range(1, ws.max_column + 1):
        celula = ws.cell(row=1, column=coluna)
        celula.font = Font(bold=True, color="FFFFFFFF")
        celula.fill = FILL_CABECALHO
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        largura = max(
            (len(str(ws.cell(row=r, column=coluna).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[get_column_letter(coluna)].width = min(max(largura + 2, 12), 45)

    ws.freeze_panes = "A2"

    for nome in ("Valor Unitário Médio (R$)", "Custo Anual Previsto (R$)"):
        coluna = cabecalhos.get(nome)
        if coluna:
            formato = "#,##0.0000" if "Unitário" in nome else "#,##0.00"
            for linha in range(2, ws.max_row + 1):
                ws.cell(row=linha, column=coluna).number_format = formato

    col_custo = cabecalhos.get("Custo Anual Previsto (R$)")
    if col_custo:
        for linha in range(2, ws.max_row + 1):
            celula = ws.cell(row=linha, column=col_custo)
            valor = celula.value
            # Antes: `if valor > 100000` estourava TypeError quando a célula
            # continha texto ou None.
            if not isinstance(valor, (int, float)):
                continue
            if valor > LIMITE_CUSTO_ALTO:
                celula.fill = FILL_VERMELHO
            elif valor > LIMITE_CUSTO_MEDIO:
                celula.fill = FILL_AMARELO
            else:
                celula.fill = FILL_VERDE

    wb.save(caminho)


def executar_analise_orcamento(
    arquivo_dispensacao: str | Path,
    arquivo_distribuicao: str | Path,
) -> Path:
    """Gera a planilha de projeção orçamentária e devolve o caminho gerado."""
    df_disp = _preparar(arquivo_dispensacao, COL_DISPENSACAO, "dispensação", "Dispensação")
    df_dist = _preparar(arquivo_distribuicao, COL_DISTRIBUICAO, "distribuição", "Distribuição")

    df = pd.concat([df_disp, df_dist], ignore_index=True)
    df = df[df["Quantidade"] > 0]
    if df.empty:
        raise ValueError("Nenhuma movimentação com quantidade válida nas planilhas enviadas.")

    # Normalização única e cacheada (antes: agrupar_medicamento por linha, sem cache).
    df["Medicamento Agrupado"] = df["Medicamento/Produto"].map(agrupar_medicamento)

    resumo = calcular_consumo(df, coluna_valor="Valor Unitário")
    if resumo.empty:
        raise ValueError("Não há datas válidas nas planilhas — verifique as colunas de data.")

    resumo["Consumo Anual Previsto"] = (resumo["Consumo Mensal Médio Corrigido"] * 12).round(2)
    resumo["Custo Anual Previsto (R$)"] = (
        resumo["Consumo Anual Previsto"] * resumo["Valor Unitário Médio (R$)"]
    ).round(2)

    resumo = resumo[
        [
            "Medicamento Agrupado",
            "Medicamento",
            "Consumo Total",
            "Dias com Falta",
            "Período Analisado (dias)",
            "Consumo Mensal Médio Corrigido",
            "Consumo Anual Previsto",
            "Valor Unitário Médio (R$)",
            "Custo Anual Previsto (R$)",
            "Amostra Insuficiente",
        ]
    ].sort_values("Custo Anual Previsto (R$)", ascending=False)

    carimbo = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    caminho = RESULTADO_DIR / f"analise_orcamento_{carimbo}.xlsx"

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        resumo.to_excel(writer, index=False, sheet_name="Orçamento")

    _formatar(caminho)

    total = float(resumo["Custo Anual Previsto (R$)"].sum())
    logger.info(
        "Análise de orçamento concluída: %d medicamentos, custo anual previsto R$ %.2f",
        len(resumo),
        total,
    )
    return caminho
